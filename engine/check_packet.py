#!/usr/bin/env python3
"""
Check that a scenario's packet (.docx) agrees with its Response Matrix (.xlsx).

    python3 check_packet.py "SCENARIO – Response Matrix.xlsx" "SCENARIO – Scenario Packet.docx"

The build generates the matrix from the packet (scenario_packet.py), so most of this is a check that the packet was
read right; what it adds is the packet's prose: every mention of the question budget, and the final-answer options.
The packet is read as it says now: pending suggestions are left out until the author accepts them. For every piece of text the game
uses (answers, Clues, Red Herrings, Major Clues, follow-ups, tripwires, the Hook, the endings, the final-answer
options), this looks for the same words in the packet. It also checks each answer's item and receptiveness
label ("Q1 [C02, clue; opinion receptive]"), the ranking order and the question budget. Curly and straight quotes
and line breaks are treated as the same. Prose that exists only in the packet (the timeline, the positions,
the hidden payoff) is not checked, and neither are the Start summary and the two final-answer questions,
which the packet words differently. The packet's 'Newspaper article' section is the one piece of the game that comes
from the packet rather than the matrix; packet_article() reads it.

Needs openpyxl and python-docx (pip3 install openpyxl python-docx).
"""
import difflib, os, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scenario_matrix import read_matrix

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("This needs python-docx: pip3 install python-docx")

NUMBER_WORDS = {"eight": 8, "ten": 10, "twenty": 20, "twenty-five": 25, "thirty": 30, "thirty-five": 35, "forty": 40,
                "forty-five": 45, "fifty": 50, "sixty": 60}


def norm(s):
    s = (s.replace("\u2018", "'").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
          .replace("\u00a0", " ").replace("\u2011", "-"))
    return " ".join(s.split())


def packet_paragraphs(path, accepted=False):
    """Every paragraph and table cell of the .docx, in order, as plain text: with its pending suggestions as if
    accepted, or with accepted=True as it says now (suggested insertions left out, suggested deletions kept)."""
    d = docx.Document(path)
    out = []
    pending = lambda x: any(a.tag == qn("w:ins") for a in x.iterancestors())
    for el in d.element.body.iter(qn("w:p")):
        t = "".join(("" if accepted and pending(x) else x.text or "") if x.tag == qn("w:t") else (x.text or "") if x.tag == qn("w:delText") else " "
                    for x in el.iter(qn("w:t"), qn("w:br"), qn("w:tab"), *([qn("w:delText")] if accepted else [])))
        if t.strip():
            out.append(norm(t))
    return out


ARTICLE_SECTION = "Newspaper article"  # the packet section ("2. Newspaper article") the game shows as the Newspaper passage


def accepted_text(p):
    """A paragraph's text as the packet says it now: pending suggestions left out. Google's .docx export carries
    them as tracked changes, a suggested insertion as text inside w:ins and a suggested deletion as w:delText."""
    out = []
    for x in p._p.iter(qn("w:t"), qn("w:delText"), qn("w:br"), qn("w:cr"), qn("w:tab"), qn("w:noBreakHyphen")):
        if x.tag == qn("w:t") and any(a.tag == qn("w:ins") for a in x.iterancestors()):
            continue
        out.append((x.text or "") if x.tag in (qn("w:t"), qn("w:delText")) else "-" if x.tag == qn("w:noBreakHyphen") else " ")
    return " ".join("".join(out).split())


def packet_article(path):
    """The packet's newspaper article, or None if the packet has no 'Newspaper article' section.

    The section is a Heading 1 named "Newspaper article" (numbered or not, such as "2. Newspaper article") and
    runs to the next Heading 1 or the end of its tab (a section break in Google's .docx export). It holds a heading (the headline), then the deck, a line starting "By " (the byline)
    and the paragraphs of the story. Two optional parts, for the printed front page only (article_pdf.py): a
    paragraph starting "Illustration:" (the illustration's caption), and a box such as "Tonight's panel": a heading
    after the byline, then bulleted lines "Name – description". The text is the packet's accepted text: pending
    suggestions are left out until the author accepts them. Returns
    ({"headline", "deck", "byline", "paragraphs", "caption", "panel"}, problems); panel is None or
    {"title", "items": [{"name", "text"}], "after": the number of story paragraphs before it}."""
    listed = lambda p: p.style.name.startswith("List") or (p._p.pPr is not None and p._p.pPr.numPr is not None)
    named = lambda t: re.fullmatch(r"(\d+\.?\s*)?" + ARTICLE_SECTION, t, re.I)
    sections, section = [], None
    for p in docx.Document(path).paragraphs:
        t = accepted_text(p)
        if p.style.name == "Heading 1":
            section = [] if named(t) else None
            if section is not None:
                sections.append(section)
        elif section is not None and t:
            section.append((p.style.name, t, listed(p)))
        if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None:
            section = None
    if not sections:
        return None, []
    if len(sections) > 1:
        return None, [f"The packet has {len(sections)} '{ARTICLE_SECTION}' sections; keep the article in one"]
    section = sections[0]
    head = next((i for i, (style, _, _) in enumerate(section) if style.startswith("Heading")), None)
    by = next((i for i, (_, t, _) in enumerate(section) if t.startswith("By ")), None)
    if head is None or by is None or by != head + 2:
        return None, [f"The '{ARTICLE_SECTION}' section must start with a heading (the headline), then one paragraph (the deck), "
                      "then the byline ('By …'), then the story"]
    body, caption, panel = [], None, None
    for style, t, is_list in section[by + 1:]:
        if t.startswith("Illustration:"):
            caption = t.removeprefix("Illustration:").strip()
        elif style.startswith("Heading"):
            if panel:
                return None, [f"The '{ARTICLE_SECTION}' section has more than one heading after the byline; it can hold one box"]
            panel = {"title": t, "items": [], "after": len(body)}
        elif is_list and panel:
            name, text = (re.split(r"\s+[–—-]\s+", t, maxsplit=1) + [""])[:2]  # "Name – description"
            panel["items"].append({"name": name, "text": text})
        else:
            body.append(t)
    if not body:
        return None, [f"The '{ARTICLE_SECTION}' section has no story after the byline"]
    if panel and not panel["items"]:
        return None, [f"The box '{panel['title']}' in the '{ARTICLE_SECTION}' section needs bulleted lines under its heading"]
    return {"headline": section[head][1], "deck": section[head + 1][1], "byline": section[by][1], "paragraphs": body,
            "caption": caption, "panel": panel}, []


def expected(s):
    """(label, text) pairs the packet must contain, and (label, regex) pairs for labels whose layout can vary."""
    W = {w["id"]: w for w in s.WITNESSES}
    texts = [("Title", s.TITLE), ("The mystery", s.PUZZLE["mystery"]), ("The solution", s.PUZZLE["solution"]),
             ("the Hook: Situation", s.SETTING["player_role"]), ("the Hook: Assignment", s.SETTING["briefing"]),
             ("the Hook: Reminder", s.SETTING["reminder"]),
             ("the Hook: CCC", s.CONTROVERSY["name"]), ("the Hook: CCC summary", s.CONTROVERSY["summary"]),
             ("the Hook: Clause", s.CONTROVERSY["clause"]), ("the Hook: Clause explainer", s.CONTROVERSY["clause_explainer"]),
             ("Map file", s.MAP)]
    if s.LOGO:
        texts.append(("Logo file", s.LOGO))
    if s.MASTHEAD:
        texts.append(("Masthead file", s.MASTHEAD))
    if s.START_IMAGE:
        texts.append(("Start image file", s.START_IMAGE))
    if s.NEWS_ART:
        texts.append(("Newspaper illustration file", s.NEWS_ART))
    for q in s.QUESTIONS:
        texts.append((f"{q['id']} question", q["text"]))
    for w in s.WITNESSES:
        n = w["name"]
        texts += [(f"{n}: background", w["background"]), (f"{n}: public card", w["public_card"]), (f"{n}: position on the CCC", w["stance"]),
                  (f"{n}: secret", w["secret"]), (f"{n}: location", w["location"]), (f"{n}: portrait", w["portrait"]),
                  (f"{n}: walk-away text", s.TRIPWIRE_RESPONSES[w["id"]])]
        for qid, a in s.RESPONSES[w["id"]].items():
            texts.append((f"{n} {qid} answer", a["text"]))
        f = s.FOLLOW_UPS[w["id"]]
        texts += [(f"{n} follow-up: receptive option", f["good"]), (f"{n} follow-up: unreceptive option", f["bad"]),
                  (f"{n} follow-up: explanation", f["why"]), (f"{n} follow-up: walk-away", f["walkout"])]
    texts += [(f"Clue {k}", c["text"]) for k, c in s.CLUES.items()]
    texts += [(f"Red Herring {k}", r["text"]) for k, r in s.RED_HERRINGS.items()]
    for k, st in s.RED_HERRING_SETS.items():
        texts += [(f"Set {k} title", st["title"]), (f"Set {k} resolution", st["resolution"])]
    for m in s.MAJOR_CLUES:
        texts += [(f"{m['id']} title", m["title"]), (f"{m['id']} text", m["text"])]
    labels = {"correct": "ending: Correct", "right_person_wrong_place": "ending: Right person, wrong place",
              "wrong": "ending: Wrong", "out_of_questions": "Submission: Out of questions"}
    texts += [(labels[k], v) for k, v in s.ENDINGS.items()]
    texts.append(("Ranking order", ", ".join(f"{i + 1}. {W[x]['name']}" for i, x in enumerate(s.BONUS_RANKING["answer"]))))

    patterns = []
    for w in s.WITNESSES:
        for qid, a in s.RESPONSES[w["id"]].items():
            kind = "clue" if a["item"] in s.CLUES else "red herring"
            flag = "receptive" if qid in w["receptive_questions"] else "unreceptive"
            patterns.append((f"{w['name']} {qid} label [{a['item']}, {kind}; opinion {flag}]", w["name"], qid,
                             rf"\b{qid}\s*\[\s*{a['item']}\s*,\s*{kind}\s*;\s*opinion\s+{flag}\s*\]"))
    # a cell of several paragraphs (split at blank lines, as the game shows them) must match consecutive packet paragraphs
    return [(l, "\n".join(norm(x) for x in re.split(r"\n\s*\n", t) if x.strip())) for l, t in texts if t], patterns


def check(matrix_path, packet_path, facts=None):
    s, errors = read_matrix(matrix_path, facts)
    if errors:
        return [f"The matrix itself has problems; fix those first ({len(errors)}, run scenario_matrix.py)"]
    paras = packet_paragraphs(packet_path, accepted=True)
    full = "\n".join(paras)
    problems = []
    texts, patterns = expected(s)
    for label, t in texts:
        if t in full:
            continue
        best = max(paras, key=lambda p: difflib.SequenceMatcher(None, t, p).ratio() if difflib.SequenceMatcher(None, t, p).real_quick_ratio() > 0.5 else 0)
        sm = difflib.SequenceMatcher(None, t.split(), best.split(), autojunk=False)
        # show only the matched stretch of the closest packet paragraph, word by word
        blocks = [b for b in sm.get_matching_blocks() if b.size]
        if blocks and sum(b.size for b in blocks) >= len(t.split()) * 0.5:
            tw, pw = t.split(), best.split()
            lo, hi = blocks[0].b, blocks[-1].b + blocks[-1].size
            diffs = []
            for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, tw, pw[lo:hi], autojunk=False).get_opcodes():
                if op != "equal":
                    diffs.append(f"matrix '{' '.join(tw[a1:a2])}' vs packet '{' '.join(pw[lo + b1:lo + b2])}'")
            problems.append(f"{label}: packet wording differs: " + "; ".join(diffs[:4]))
        else:
            problems.append(f"{label}: not found in the packet: \"{t[:90]}{'…' if len(t) > 90 else ''}\"")
    for label, name, qid, rx in patterns:
        if not re.search(rx, full, re.I):
            problems.append(f"{label}: the packet does not label this answer the same way")

    # Final-answer options: the packet must list every option the game offers, either in one paragraph or spread over
    # the final-answer table (the paragraphs from "Final-answer form" to "The aftermath", for example one per row).
    # The check stays inside that one place, since the names also appear elsewhere in the packet.
    options = [("person", o) for o in s.PUZZLE["who_options"]] + [("place", o) for o in s.PUZZLE["where_options"]]
    home = max(paras, key=lambda p: sum(norm(o) in p for _, o in options))
    start = next((i for i, p in enumerate(paras) if p.startswith(norm("Final-answer form"))), None)
    table = []
    if start is not None:
        for p in paras[start + 1:start + 61]:
            if p.startswith(norm("The aftermath")):
                break
            table.append(p)
    for kind, o in options:
        if norm(o) not in home and not any(norm(o) in p for p in table):
            problems.append(f"Final-answer form: the packet's list of offered options is missing the {kind} \"{o}\"")
    fa = s.PUZZLE["final_answer"]
    for kind in ("who", "where"):
        if norm(fa[kind]) not in home and not any(norm(fa[kind]) in p for p in table):
            problems.append(f"Final-answer form: the correct {kind} (\"{fa[kind]}\") is not given with the options")

    # Question budget (a framework fact): every "<number> questions" in the packet must be the budget (or the eight questions)
    for m in re.finditer(r"\b(\d+|[a-z]+(?:-[a-z]+)?)\s+questions\b", full, re.I):
        word = m.group(1).lower()
        n = int(word) if word.isdigit() else NUMBER_WORDS.get(word)
        if n is not None and n not in (s.QUESTION_BUDGET, len(s.QUESTIONS)):
            ctx = full[max(0, m.start() - 40):m.end() + 10].replace("\n", " ")
            problems.append(f"Question budget: the packet says '{m.group(0)}' but the framework's budget is {s.QUESTION_BUDGET} (…{ctx}…)")
    return problems


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    problems = check(sys.argv[1], sys.argv[2])
    print("PACKET DISAGREES WITH THE MATRIX:" if problems else "OK: the packet matches the matrix")
    for p in problems:
        print(" -", p)
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
