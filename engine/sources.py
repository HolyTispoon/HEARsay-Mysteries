#!/usr/bin/env python3
"""
The sources of truth, the repository's copies of them, and the workbooks generated from them.

The sources of truth are Google Docs in Drive, outside this repository, and only the author edits them:
  - the framework Google Doc "HEARsay mysteries": the framework, and in five of its tabs the game's own text and the
    framework's numbers (game_text.py);
  - for each scenario, its Scenario Packet (a Google Doc): the whole scenario (scenario_packet.py).
Their addresses are in engine/sources.json and scenarios/<id>/sources.json.

The repository keeps a copy of each Doc (engine/HEARsay mysteries.docx and, in each scenario folder, the packet
.docx). Nobody edits a copy. update_copies() replaces each copy with a download of its Google Doc (File → Download →
Microsoft Word, or the Drive connector's export) when the download's content differs, after checking every download.
Then it writes the workbooks the build reads, generated from the copies: engine/HEARsay game text.xlsx from the
framework and, in each scenario folder, the Response Matrix .xlsx from the packet. Nobody edits those either; the
build rewrites them whenever their content would change.

    python3 engine/sources.py [DOWNLOADS_FOLDER]    # lists the sources, and which downloads would be used
"""
import glob, json, os, shutil, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCENARIOS = os.path.join(ROOT, "scenarios")
sys.path.insert(0, HERE)
PACKET, MATRIX = "Scenario Packet", "Response Matrix"  # in the file names: "<Title> – Scenario Packet.docx"


def _load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def packet_copy(sid):
    """The copy of a scenario's packet."""
    folder = os.path.join(SCENARIOS, sid)
    found = [p for p in glob.glob(os.path.join(glob.escape(folder), "*.docx"))
             if PACKET.lower() in os.path.basename(p).lower() and not os.path.basename(p).startswith("~$")]
    if len(found) != 1:
        sys.exit(f"Expected one packet copy ('… – {PACKET}.docx') in scenarios/{sid}/, found {len(found)}")
    return found[0]


def matrix_path(sid):
    """The Response Matrix workbook generated from a scenario's packet: the packet's name with 'Response Matrix'."""
    packet = packet_copy(sid)
    name = os.path.basename(packet)
    i = name.lower().index(PACKET.lower())
    return os.path.join(os.path.dirname(packet), name[:i] + MATRIX + ".xlsx")


def copies(ids):
    """(label, copy path, Google address) for the framework and each scenario's packet, in build order."""
    from game_text import FRAMEWORK
    out = [("framework", FRAMEWORK, _load(os.path.join(HERE, "sources.json"))["framework"])]
    for sid in ids:
        urls = _load(os.path.join(SCENARIOS, sid, "sources.json"))
        out.append((f"{sid} packet", packet_copy(sid), urls.get("packet") or "(not in Google Drive yet)"))
    return out


def find_download(folder, copy_path):
    """The newest download in folder named like the copy ('The Commons' Comma – Scenario Packet (1).docx' too)."""
    stem, ext = os.path.splitext(os.path.basename(copy_path))
    found = [p for p in glob.glob(os.path.join(glob.escape(folder), glob.escape(stem) + "*" + ext))
             if not os.path.basename(p).startswith("~$")]
    return max(found, key=os.path.getmtime) if found else None


def content(path):
    """What the build reads from a file: every sheet's cell values, or every paragraph of a document."""
    if path.endswith(".docx"):
        from check_packet import packet_paragraphs
        return packet_paragraphs(path)
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb:
        rows = [[None if v == "" else v for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r[:max([i + 1 for i, v in enumerate(r) if v is not None], default=0)] for r in rows]
        while rows and not rows[-1]:
            rows.pop()
        out[ws.title] = rows
    return out


def generate(use, ids, folder):
    """Write the workbooks generated from the Docs in use ({label: path}) into folder. Returns ({label: workbook
    path}, problems)."""
    from game_text import framework_workbook, DEFAULT as GAME_TEXT
    from scenario_matrix import write_workbook
    from scenario_packet import read_packet
    out, problems = {}, []
    sheets, errors = framework_workbook(use["framework"])
    problems += [f"framework: {e}" for e in errors]
    if not errors:
        out["game text"] = os.path.join(folder, os.path.basename(GAME_TEXT))
        write_workbook(sheets, out["game text"])
    for sid in ids:
        sheets, errors = read_packet(use[f"{sid} packet"])
        problems += [f"{sid} packet: {e}" for e in errors]
        if not errors:
            out[f"{sid} matrix"] = os.path.join(folder, os.path.basename(matrix_path(sid)))
            write_workbook(sheets, out[f"{sid} matrix"])
    return out, problems


def check_downloads(use, ids, workbooks):
    """Problems with the game text and each scenario, read from the generated workbooks."""
    from game_text import read_game_text
    from scenario_matrix import read_matrix
    from check_packet import check as check_packet
    g, errors = read_game_text(workbooks["game text"], use["framework"])
    if errors:
        return [f"game text: {e}" for e in errors]
    problems = []
    for sid in ids:
        _, merr = read_matrix(workbooks[f"{sid} matrix"], g.facts)
        problems += [f"{sid} matrix (from the packet): {e}" for e in merr]
        if not merr:
            problems += [f"{sid} packet: {p}" for p in check_packet(workbooks[f"{sid} matrix"], use[f"{sid} packet"], g.facts)]
    return problems


def update_copies(folder, ids):
    """Replace each copy with its download from folder when the content differs, then rewrite each generated
    workbook whose content differs. Stops the build, changing nothing, if a Doc has problems. Returns the labels of
    the files that changed."""
    from game_text import DEFAULT as GAME_TEXT
    chosen = [(label, path, url, find_download(folder, path) if folder else None) for label, path, url in copies(ids)]
    if folder and not any(d for *_, d in chosen):
        sys.exit(f"No downloads in {folder}. Download the Google Docs as .docx (File → Download → Microsoft Word), "
                 "named as they are in Drive, into that folder.")
    use = {label: download or path for label, path, _, download in chosen}
    scratch = tempfile.mkdtemp()
    workbooks, problems = generate(use, ids, scratch)
    if not problems:
        problems = check_downloads(use, ids, workbooks)
    if problems:
        sys.exit("The Google Docs have problems; fix them in the Docs, download again, and rebuild. Nothing was "
                 "copied. (A pending suggestion is not read until it is accepted.)\n" + "\n".join(" - " + p for p in problems))
    changed = []
    for label, path, url, download in chosen:
        rel = os.path.relpath(path, ROOT)
        if not download:
            print(f"WARNING: {label}: no download, so the build uses the copy {rel}, which may be behind {url}")
        elif content(download) == content(path):
            print(f"{label}: the copy already matches {os.path.basename(download)}")
        else:
            shutil.copyfile(download, path)
            changed.append(label)
            print(f"{label}: updated the copy {rel} from {os.path.basename(download)}")
    targets = [("game text", GAME_TEXT)] + [(f"{sid} matrix", matrix_path(sid)) for sid in ids]
    for label, path in targets:
        rel = os.path.relpath(path, ROOT)
        if os.path.exists(path) and content(workbooks[label]) == content(path):
            continue
        shutil.copyfile(workbooks[label], path)
        changed.append(label)
        print(f"{label}: regenerated {rel}")
    shutil.rmtree(scratch, ignore_errors=True)
    if changed:
        print("Commit the updated copies and workbooks with the rebuilt docs/.")
    return changed


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else None
    ids = sorted(n for n in os.listdir(SCENARIOS) if os.path.isdir(os.path.join(SCENARIOS, n)) and not n.startswith("."))
    for label, path, url in copies(ids):
        d = find_download(folder, path) if folder else None
        print(f"{label}: {url}\n    copy: {os.path.relpath(path, ROOT)}" + (f"\n    download: {d}" if d else ""))
    for sid in ids:
        print(f"{sid} matrix: generated from the packet\n    {os.path.relpath(matrix_path(sid), ROOT)}")


if __name__ == "__main__":
    main()
