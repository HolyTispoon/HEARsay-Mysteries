#!/usr/bin/env python3
"""
Read the game's own text and the framework's numbers from the framework's 'HEARsay mysteries.docx'.

    python3 game_text.py ["HEARsay game text.xlsx" ["HEARsay mysteries.docx"]]

The source of truth is the framework Google Doc "HEARsay mysteries" (its address is in sources.json). The .docx next
to this script is the repository's copy of it, which only the build updates (see sources.py), and the .xlsx is
generated from it by the build (framework_workbook()); nobody edits either.

The framework Doc's tabs (each exported as a Title paragraph with the tab's name, up to the next tab):
  - 'Rules of the game': a heading (rules.title), then one paragraph per rule (rules.1, rules.2, …);
  - 'HEAR Framework introduction': a heading (hear.title), a paragraph (hear.intro), then for H, E, A and R a
    heading (hear.H.move …) and a paragraph (hear.H.text …), then a last paragraph (hear.unreceptive);
  - 'Game facts': a table with the columns Key, Fact, Value, Framework wording and Framework section, one row per
    number of the framework (check_framework.py checks the wording against the rest of the Doc);
  - 'Game text': a table with the columns Key, Screen, Text and Placeholders, for the game's sentences that
    restate the framework;
  - 'Screens and labels': a table of the game's screens (Engine passage, Framework name, Framework section, Notes)
    and a table of the rest of the game's text, its headings, buttons and labels (Key, Engine passage, Text,
    Framework section, Check, Placeholders). The build writes them to the workbook's Passages and Text sheets.
Bold words in the Doc become **double asterisks**, which the game shows in bold. Pending suggestions are not read.

read_game_text() returns (game_text, errors). game_text.text maps each key to its text, game_text.facts maps each
fact key to its value, and game_text.rows / .passages / .fact_rows keep the rows for check_framework.py.
Needs openpyxl and python-docx (pip3 install openpyxl python-docx).
"""
import os, re, sys
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scenario_matrix import _sheet, _keyed

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("This needs openpyxl: pip3 install openpyxl")

DEFAULT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HEARsay game text.xlsx")  # generated from the framework Doc
FRAMEWORK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HEARsay mysteries.docx")  # the copy of the Google Doc
# The framework tabs the game reads
RULES_TAB, HEAR_TAB, FACTS_TAB, TEXT_TAB = "Rules of the game", "HEAR Framework introduction", "Game facts", "Game text"
SCREENS_TAB = "Screens and labels"
FACT_COLUMNS = ["Key", "Fact", "Value", "Framework wording", "Framework section"]
TEXT_COLUMNS = ["Key", "Screen", "Text", "Placeholders"]
# The 'Screens and labels' tab's two tables, which become the workbook's Passages and Text sheets
SCREEN_COLUMNS = ["Engine passage", "Framework name", "Framework section", "Notes"]
LABEL_COLUMNS = ["Key", "Engine passage", "Text", "Framework section", "Check", "Placeholders"]
WORKBOOK_README = ("Generated from the framework Google Doc 'HEARsay mysteries' (its 'Screens and labels' tab) by the "
                   "build; don't edit it. Change the Doc and rebuild. game_text.py says how the build reads the Doc.")
# Keys that come from the rules and HEAR tabs, so not from the labels table
DOC_KEYS = re.compile(r"^(rules\.(title|\d+)|hear\.(title|intro|unreceptive|[HEAR]\.(move|text)))$")
# Placeholders every text can use; the engine fills them from the scenario and the facts.
GLOBAL_PLACEHOLDERS = {"budget", "n_witnesses", "n_questions", "n_major", "rank_exact", "rank_near", "ccc", "title"}
CHECKS = {"verbatim", "paraphrase", "label"}
# Facts the build and the engine rely on.
REQUIRED_FACTS = ["witnesses", "questions", "question_budget", "tripwires_per_witness", "clues_per_witness",
                  "red_herrings_per_witness", "clues", "givers_per_clue", "red_herrings", "givers_per_red_herring",
                  "rh_set_min", "rh_set_max", "major_clues", "clues_per_major", "followups_per_witness", "submissions",
                  "rank_exact", "rank_near"]


def _paragraph_text(p):
    """A paragraph's text with its bold runs in **double asterisks** (bold that comes from a heading style is left
    alone)."""
    out = []
    for r in p.runs:
        t = r.text
        out.append(f"**{t.strip()}**" if r.bold and t.strip() and "**" not in t else t)
    t = " ".join("".join(out).split())
    return re.sub(r"\*\*\s*\*\*", " ", t).strip()  # two bold runs side by side make one bold phrase


def framework_blocks(path):
    """The framework export in order, as (tab, style, text) for each non-empty paragraph and (tab, "Table", rows)
    for each table, rows being lists of cell texts. Paragraphs before the first tab title have tab None.

    Google's .docx export has marked tabs in two ways: with a Title paragraph holding the tab's name, or (since
    October 2026, seen in the packets) only with a section break at the end of each tab. When one of the tabs the game
    reads has no name in the export, the section that has its layout is taken as that tab (_shape_tab)."""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    blocks, sections, tab, section = [], [], None, 0
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            p = Paragraph(el, d)
            t = " ".join(p.text.split())
            if p.style.name == "Title":
                tab = t
            elif t:
                blocks.append((tab, p.style.name, _paragraph_text(p)))
                sections.append((section, el.pPr is not None and el.pPr.numPr is not None))
            if el.pPr is not None and el.pPr.find(qn("w:sectPr")) is not None:
                section += 1
        elif el.tag.endswith("}tbl"):
            rows = [[" ".join(_paragraph_text(p) for p in c.paragraphs).strip() for c in r.cells] for r in Table(el, d).rows]
            blocks.append((tab, "Table", rows))
            sections.append((section, False))
    named = {b[0] for b in blocks}
    for name in (RULES_TAB, HEAR_TAB, FACTS_TAB, TEXT_TAB, SCREENS_TAB):
        if name in named:
            continue
        by_section = {}
        for i, (n, _) in enumerate(sections):
            by_section.setdefault(n, []).append(i)
        found = [ix for ix in by_section.values() if _shape_tab([blocks[i] for i in ix], [sections[i][1] for i in ix]) == name]
        if len(found) == 1:
            for i in found[0]:
                blocks[i] = (name,) + blocks[i][1:]
    return blocks


def _shape_tab(blocks, listed):
    """Which of the tabs the game reads a section's blocks are laid out as, if any. listed: whether each block is a
    list item."""
    styles = [style for _, style, _ in blocks]
    heading = [s.startswith("Heading") for s in styles]
    for _, style, rows in blocks:
        if style == "Table" and rows:
            head = [h.replace("*", "").strip() for h in rows[0]]
            if head[:len(FACT_COLUMNS)] == FACT_COLUMNS:
                return FACTS_TAB
            if head[:len(TEXT_COLUMNS)] == TEXT_COLUMNS:
                return TEXT_TAB
            if head[:len(SCREEN_COLUMNS)] == SCREEN_COLUMNS or head[:len(LABEL_COLUMNS)] == LABEL_COLUMNS:
                return SCREENS_TAB
    if "Table" in styles:
        return None
    if heading == [True, False] + [True, False] * 4 + [False] and \
            "".join(blocks[2 + 2 * i][2].replace("*", "").strip()[:1].upper() for i in range(4)) == "HEAR":
        return HEAR_TAB
    if len(blocks) > 2 and heading[0] and not any(heading[1:]) and all(listed[1:]):
        return RULES_TAB
    return None


def _table(blocks, tab, columns, errors):
    """The rows of the one table in a tab, as dicts keyed by column, with their row numbers in '_row'."""
    tables = [rows for t, style, rows in blocks if t == tab and style == "Table"]
    if len(tables) != 1:
        errors.append(f"The framework's '{tab}' tab must hold one table, not {len(tables)}")
        return []
    head = [h.replace("*", "").strip() for h in tables[0][0]]
    if head[:len(columns)] != columns:
        errors.append(f"The table in the framework's '{tab}' tab must have the columns {', '.join(columns)}, "
                      f"not {', '.join(head)}")
        return []
    plain = {"Key", "Value", "Screen", "Placeholders", "Framework section"}  # cells that are names or numbers, never bold
    return [dict({c: (r[i].replace("*", "").strip() if c in plain else r[i]) for i, c in enumerate(columns)}, _row=n)
            for n, r in enumerate(tables[0][1:], 2) if any(x.strip() for x in r)]


def _placeholders(where, text, allowed, errors):
    for ph in re.findall(r"\{(\w+)\}", text):
        if ph not in allowed:
            errors.append(f"{where}: {{{ph}}} is not a placeholder this text can use (allowed: {', '.join(sorted(allowed))})")


def framework_text(path, engine_passages):
    """From the framework's tabs: ({key: text}, fact rows, errors)."""
    if not os.path.exists(path):
        return {}, [], [f"No copy of the framework Doc ({os.path.basename(path)}): download the Google Doc as .docx "
                        "into the downloads folder and build with --downloads"]
    blocks = framework_blocks(path)
    tab = lambda name: [(style, t) for b, style, t in blocks if b == name and style != "Table"]
    text, errors = {}, []
    heading = lambda style: style.startswith("Heading")

    rules = tab(RULES_TAB)
    if len(rules) < 2 or not heading(rules[0][0]) or any(heading(s) for s, _ in rules[1:]):
        errors.append(f"The framework's '{RULES_TAB}' tab must hold a heading, then one paragraph per rule")
    else:
        text["rules.title"] = rules[0][1].replace("*", "")
        text.update((f"rules.{i}", t) for i, (_, t) in enumerate(rules[1:], 1))

    hear = tab(HEAR_TAB)
    if [heading(s) for s, _ in hear] != [True, False] + [True, False] * 4 + [False]:
        errors.append(f"The framework's '{HEAR_TAB}' tab must hold a heading and a paragraph, then a heading and a "
                      "paragraph for each of H, E, A and R, then one last paragraph (unreceptive language)")
    else:
        text["hear.title"], text["hear.intro"], text["hear.unreceptive"] = hear[0][1].replace("*", ""), hear[1][1], hear[-1][1]
        for i, L in enumerate("HEAR"):
            move = hear[2 + 2 * i][1].replace("*", "")
            text[f"hear.{L}.move"], text[f"hear.{L}.text"] = move, hear[3 + 2 * i][1]
            if not move.upper().startswith(L):
                errors.append(f"The framework's '{HEAR_TAB}' tab: the {L} heading must start with {L}, not {move!r}")
    for k, t in text.items():
        _placeholders(f"The framework's {k}", t, GLOBAL_PLACEHOLDERS, errors)

    for r in _table(blocks, TEXT_TAB, TEXT_COLUMNS, errors):
        k, where = r["Key"], f"The framework's '{TEXT_TAB}' row {r['_row']} ({r['Key'] or 'no key'})"
        if not k:
            errors.append(f"{where}: Key is blank"); continue
        if k in text or DOC_KEYS.match(k):
            errors.append(f"{where}: key appears twice in the framework"); continue
        if r["Screen"] not in engine_passages:
            errors.append(f"{where}: screen {r['Screen']!r} is not in the '{SCREENS_TAB}' tab's table of screens")
        _placeholders(where, r["Text"], GLOBAL_PLACEHOLDERS | set(r["Placeholders"].replace(",", " ").split()), errors)
        text[k] = r["Text"]

    fact_rows = _table(blocks, FACTS_TAB, FACT_COLUMNS, errors)
    return text, fact_rows, errors


def framework_workbook(framework):
    """The game text workbook's sheets from the framework's 'Screens and labels' tab: (sheets, errors)."""
    if not os.path.exists(framework):
        return {}, [f"No copy of the framework Doc ({os.path.basename(framework)})"]
    tables = [rows for _, style, rows in framework_blocks(framework) if style == "Table" and rows]
    sheets, errors = {"Read me": [["Sheet", "What it holds"], ["", WORKBOOK_README]]}, []
    for name, columns, what in (("Passages", SCREEN_COLUMNS, "screens"), ("Text", LABEL_COLUMNS, "labels")):
        found = [t for t in tables if [h.replace("*", "").strip() for h in t[0]] == columns]
        if len(found) != 1:
            errors.append(f"The framework needs one table of {what} (in its '{SCREENS_TAB}' tab) with the columns "
                          f"{', '.join(columns)}; found {len(found)}")
            continue
        text = columns.index("Text") if "Text" in columns else None
        # only the Text column can be bold; elsewhere a bold cell is just formatting
        sheets[name] = [columns] + [[c if i == text else c.replace("*", "").strip() for i, c in enumerate(r[:len(columns)])]
                                    for r in found[0][1:] if any(x.strip() for x in r)]
    return sheets, errors


def read_game_text(path=DEFAULT, framework=FRAMEWORK):
    errors = []
    wb = load_workbook(path, data_only=True)

    passages = _sheet(wb, "Passages", ["Engine passage", "Framework name", "Framework section", "Notes"], errors)
    engine_passages = {p["Engine passage"] for p in passages}

    rows = _sheet(wb, "Text", ["Key", "Engine passage", "Text", "Framework section", "Check", "Placeholders"], errors)
    text = {}
    for r in rows:
        k, where = r["Key"], f"Labels table row {r['_row']} ({r['Key'] or 'no key'})"
        if not k:
            errors.append(f"Labels table row {r['_row']}: Key is blank"); continue
        if k in text:
            errors.append(f"{where}: key appears twice"); continue
        if r["Engine passage"] not in engine_passages:
            errors.append(f"{where}: engine passage {r['Engine passage']!r} is not in the table of screens")
        check = r["Check"].lower()
        if check not in CHECKS:
            errors.append(f"{where}: Check must be verbatim, paraphrase or label, not {r['Check']!r}")
        if check in ("verbatim", "paraphrase") and not r["Framework section"]:
            errors.append(f"{where}: a {check} row needs a Framework section")
        _placeholders(where, r["Text"], GLOBAL_PLACEHOLDERS | set(r["Placeholders"].split()), errors)
        if check == "verbatim" and "{" in r["Text"]:
            errors.append(f"{where}: a verbatim row cannot contain placeholders")
        text[k] = r["Text"]

    doc_text, doc_facts, doc_errors = framework_text(framework, engine_passages)
    errors += doc_errors
    for r in rows:
        if r["Key"] in doc_text or DOC_KEYS.match(r["Key"]):
            errors.append(f"Labels table row {r['_row']} ({r['Key']}): this text comes from another tab of the framework; delete the row")
    text.update(doc_text)
    fact_rows = _keyed(doc_facts, "Key", f"The framework's '{FACTS_TAB}'", errors)
    facts = {}
    for k, r in fact_rows.items():
        try:
            v = float(r["Value"])
            facts[k] = int(v) if v.is_integer() else v
        except ValueError:
            errors.append(f"The framework's fact {k}: Value must be a number, not {r['Value']!r}")
    for k in REQUIRED_FACTS:
        if k not in facts:
            errors.append(f"The framework's '{FACTS_TAB}': '{k}' is missing")
    return SimpleNamespace(text=text, facts=facts, rows=rows, passages=passages, fact_rows=list(fact_rows.values())), errors


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    g, errors = read_game_text(path, sys.argv[2] if len(sys.argv) > 2 else FRAMEWORK)
    print(f"{len(g.text)} texts, {len(g.passages)} passages, {len(g.facts)} facts")
    print("PROBLEMS:" if errors else "OK: the game text passes every check")
    for e in errors:
        print(" -", e)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
