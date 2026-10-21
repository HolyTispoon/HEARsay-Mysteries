#!/usr/bin/env python3
"""
Build HEARsay scenarios into playable Twine (SugarCube 2) games, and the site's landing page.

    python3 engine/build.py [SCENARIO_ID ...] [--downloads FOLDER] [--no-packet-check]
                            [--framework "HEARsay mysteries.docx" | --no-framework-check]

With no Scenario ID it builds every scenario in scenarios/. Each scenario is a folder scenarios/<Scenario ID>/
holding the copies of its Response Matrix (.xlsx) and Scenario Packet (.docx), sources.json (their Google
addresses) and assets/ (the portraits and the campus map).

The sources of truth are Google Docs and Sheets in Drive that only the author edits (see sources.py):
  - the framework Google Doc ("HEARsay mysteries"): the game's structure and rules;
  - the game text Google Sheet ("HEARsay game text"): the game's own text and the framework's numbers, checked
    against the framework;
  - each scenario's Response Matrix (Google Sheet): the scenario, checked against the framework's numbers;
  - each scenario's packet (Google Doc): the scenario written up for people, checked against the matrix.
The build reads the repository's copies of the Sheets and the packets.

Steps:
  0. With --downloads, a folder holding downloads of the Google files (Sheets as .xlsx, Docs as .docx, named as in
     Drive): checks them all, then replaces each copy whose content differs. A download with problems stops the
     build and nothing is copied. Copies with no download are used as they are, with a warning.
  1. Reads the game text and checks it against a download of the framework (check_framework.py): --framework, or
     the newest 'HEARsay mysteries' .docx/.txt/.pdf in the downloads folder, the repository folder or the folder
     above it. If there is none the check is skipped with a warning.
  2. Checks that every text the engine asks for is in the game text workbook.
  3. For each scenario: reads the matrix and stops if it breaks a framework rule or disagrees with itself
     (scenario_matrix.py), and checks that the packet says the same things as the matrix (check_packet.py).
  4. Writes docs/<Scenario ID>/source/story.twee (Twee 3), with the scenario and the game text written into the
     engine as `setup.scenario` and `setup.gameText`, and compiles it with Tweego into docs/<Scenario ID>/index.html.
     The newspaper article, if the packet has a 'Newspaper article' tab, comes from the packet (packet_article).
  5. Copies the scenario's portraits, map, logo and masthead (if it has them), and the site icons (ICONS), into
     docs/<Scenario ID>/assets/; a logo or masthead not in the scenario's assets/ is looked for in SHARED_ASSETS;
     head.html links the icons from the page's <head>.
  6. Writes the landing page, docs/index.html, from landing.html and every built scenario.
Every check stops the build and lists each problem. GitHub Pages serves docs/.

Needs openpyxl and python-docx (pip3 install openpyxl python-docx), and pypdf for a PDF framework export.

Tweego: https://www.motoslave.net/tweego/ . Put `tweego` on your PATH, or set TWEEGO=/path/to/tweego.
Tweego must have the SugarCube 2.37.3 story format (set TWEEGO_PATH to a folder holding sugarcube-2/ if the
bundled one is older).
"""
import argparse, glob, json, os, re, shutil, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)  # the repository
SCENARIOS = os.path.join(ROOT, "scenarios")
DOCS = os.path.join(ROOT, "docs")  # the published site
sys.path.insert(0, HERE)
from scenario_matrix import read_matrix, PASSAGE_PARTS
from check_packet import check as check_packet, norm, packet_article, ARTICLE_TAB
from game_text import read_game_text, DEFAULT as GAME_TEXT
from sources import update_copies
from check_framework import check as check_framework, find_export

ENGINE = os.path.join(HERE, "hearsay_engine.twee")
# Appended to the <head> of index.html; it links the site icons, which are the same for every scenario
HEAD = os.path.join(HERE, "head.html")
LANDING = os.path.join(HERE, "landing.html")  # the site's landing page; {{SCENARIOS}} becomes the list of games
ICONS = ("favicon.svg", "favicon-32.png", "apple-touch-icon.png")
# Images more than one scenario may use, such as a university's logo or its newspaper's masthead. A scenario's own assets/ come first.
SHARED_ASSETS = os.path.join(HERE, "shared-assets")
# Text keys the engine builds at run time rather than naming outright
DYNAMIC_KEYS = re.compile(r"^(rules\.\d+|hear\.[HEAR]\.(move|text)|result\.\w+|ending\.verdict\.\w+)$")


def load_game_text(framework, skip, downloads):
    g, errors = read_game_text(GAME_TEXT)
    if errors:
        sys.exit("The game text workbook has problems; fix them before building:\n" + "\n".join(" - " + e for e in errors))
    if skip:
        print("Skipping the framework check (--no-framework-check).")
        return g
    note = None
    if not framework and downloads:
        framework, note = find_export(downloads)
    if not framework:
        framework, note = find_export(ROOT)
    if not framework:
        framework, note = find_export(os.path.dirname(ROOT))  # in Drive: the folder holding the Google Doc
    if not framework:
        print("WARNING: no export of the framework ('HEARsay mysteries' .docx/.txt/.pdf) found, so the game text "
              "was not checked against it. Export the Google Doc next to the repository folder, or pass --framework.")
        return g
    problems, notes = check_framework(GAME_TEXT, framework)
    # the matrix's passage names must be framework names too
    from check_framework import framework_text, passage_list
    listed = [n.lower() for n in (passage_list(*framework_text(framework)) or [])]
    for pas in sorted({p for p, _ in PASSAGE_PARTS}):
        if pas.lower() not in listed:
            problems.append(f"The matrix files scenario text under '{pas}', which is not in the framework's passage list")
    if problems:
        sys.exit(f"The game text disagrees with the framework ({os.path.basename(framework)}); make them match before building:\n"
                 + "\n".join(" - " + p for p in problems))
    print(f"Framework OK: the game text matches {os.path.basename(framework)}"
          + (f" ({len(notes)} engine choices the framework does not describe; run check_framework.py to list them)" if notes else ""))
    if note:
        print("WARNING: " + note)
    return g


def check_engine_keys(g):
    with open(ENGINE, encoding="utf-8") as f:
        engine = f.read()
    # literal keys only; keys built at run time (T("hear." + L + ".move")) are covered by DYNAMIC_KEYS
    used = {k for k in re.findall(r"""\b[TR]\(\s*["'`]([\w.]+)["'`]""", engine) if not k.endswith(".")}
    missing = sorted(k for k in used if k not in g.text)
    if missing:
        sys.exit("The engine uses text that is not in the game text workbook: " + ", ".join(missing))
    unused = sorted(k for k in g.text if k not in used and not DYNAMIC_KEYS.match(k))
    if unused:
        print("Note: the game text workbook has text the engine never shows: " + ", ".join(unused))


def load(matrix_path, facts):
    d, errors = read_matrix(matrix_path, facts)
    if errors:
        sys.exit("The matrix has problems; fix them before building:\n" + "\n".join(" - " + e for e in errors))
    print(f"Matrix OK: {len(d.WITNESSES)} witnesses, {len(d.CLUES)} clues, {len(d.RED_HERRINGS)} red herrings, budget {d.QUESTION_BUDGET}")
    return d


def find_packet(matrix_path):
    found = [p for p in glob.glob(os.path.join(os.path.dirname(os.path.abspath(matrix_path)), "*.docx"))
             if "packet" in os.path.basename(p).lower() and not os.path.basename(p).startswith("~$")]
    if len(found) != 1:
        sys.exit(f"Expected one packet .docx next to the matrix, found {len(found)}. The build reads the packet even with --no-packet-check, for its newspaper article.")
    return found[0]


def article_json(d, article):
    """The packet's article for the game. A paragraph that quotes the matrix's Clause is marked as a quote; letters
    and digits are compared, so quotation marks and a "(,)" don't matter."""
    if not article:
        return None
    bare = lambda t: re.sub(r"\W+", "", t).lower()
    clause = bare(d.CONTROVERSY["clause"])
    quote = lambda t: bool(clause) and bare(t) == clause
    return dict(article, paragraphs=[{"text": t, "quote": quote(t)} for t in article["paragraphs"]])


def scenario_json(d, article=None, asset_dir="assets"):
    p = d.PUZZLE
    for key in ("who", "where"):
        opts = p[f"{key}_options"]
        if p["final_answer"][key] not in opts:
            sys.exit(f"PUZZLE['final_answer']['{key}'] is not one of PUZZLE['{key}_options']")
    # Only what the game needs. Designer-only fields (secrets, receptive flags, tripwire logic) stay out.
    return {
        "id": d.SCENARIO_ID,
        "title": d.TITLE,
        "pitch": d.PITCH or p["mystery"],
        "budget": d.QUESTION_BUDGET,
        "followup_bonus": d.FOLLOWUP_BONUS,
        "asset_dir": asset_dir,
        "map": d.MAP,
        "logo": d.LOGO,
        "masthead": d.MASTHEAD,
        "article": article_json(d, article),
        "rank_exact": d.FACTS["rank_exact"],
        "rank_near": d.FACTS["rank_near"],
        "puzzle": {
            "mystery": p["mystery"],
            "solution": p["solution"],
            "who_label": p["who_label"],
            "where_label": p["where_label"],
            "final_answer": p["final_answer"],
            "who_options": p["who_options"],
            "where_options": p["where_options"],
        },
        "controversy": {k: d.CONTROVERSY[k] for k in ("name", "summary", "clause", "clause_explainer")},
        "setting": {"player_role": d.SETTING["player_role"], "briefing": d.SETTING["briefing"]},
        "questions": [{"id": q["id"], "text": q["text"]} for q in d.QUESTIONS],
        "witnesses": [{k: w[k] for k in ("id", "name", "portrait", "role", "public_card", "location", "tripwire")}
                      for w in d.WITNESSES],
        "responses": {wid: {qid: {"item": a["item"], "text": a["text"]} for qid, a in qs.items()}
                      for wid, qs in d.RESPONSES.items()},
        "tripwire_responses": d.TRIPWIRE_RESPONSES,
        "follow_ups": d.FOLLOW_UPS,
        "clues": {k: v["text"] for k, v in d.CLUES.items()},
        "red_herrings": {k: v["text"] for k, v in d.RED_HERRINGS.items()},
        "red_herring_sets": d.RED_HERRING_SETS,
        "major_clues": [{k: m[k] for k in ("id", "letter", "title", "text")} for m in d.MAJOR_CLUES],
        "endings": d.ENDINGS,
        "ranking_answer": d.BONUS_RANKING["answer"],
    }


def find_file(root, name):
    for dp, _, files in os.walk(root):
        if name in files:
            return os.path.join(dp, name)
    sys.exit(f"Asset {name} not found under {root}")


def same(a, b):
    return os.path.exists(b) and os.path.samefile(a, b)


def copy(src, dst):
    if same(src, dst):
        return
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    shutil.copyfile(src, dst)


def scenario_files(sid):
    """The matrix and the scenario folder of a Scenario ID."""
    folder = os.path.join(SCENARIOS, sid)
    if not os.path.isdir(folder):
        sys.exit(f"No scenario folder {os.path.relpath(folder, ROOT)}. Scenarios: {', '.join(all_scenarios()) or 'none'}")
    found = [p for p in glob.glob(os.path.join(folder, "*.xlsx")) if not os.path.basename(p).startswith("~$")]
    if len(found) != 1:
        sys.exit(f"Expected one Response Matrix .xlsx in {os.path.relpath(folder, ROOT)}, found {len(found)}")
    return found[0], folder


def all_scenarios():
    return sorted(n for n in os.listdir(SCENARIOS) if os.path.isdir(os.path.join(SCENARIOS, n)) and not n.startswith("."))


def build_scenario(sid, g, no_packet_check):
    matrix, folder = scenario_files(sid)
    print(f"\n== {sid}")
    d = load(matrix, g.facts)
    if d.SCENARIO_ID != sid:
        sys.exit(f"The matrix's Scenario ID is '{d.SCENARIO_ID}', but its folder is scenarios/{sid}/. Make them the same.")
    packet = find_packet(matrix)
    if no_packet_check:
        print("Skipping the packet check (--no-packet-check).")
    else:
        problems = check_packet(matrix, packet)
        if problems:
            sys.exit(f"The packet ({os.path.basename(packet)}) disagrees with the matrix; make them match before building:\n"
                     + "\n".join(" - " + p for p in problems))
        print(f"Packet OK: {os.path.basename(packet)} matches the matrix")
    article, problems = packet_article(packet)
    if problems:
        sys.exit(f"The packet's newspaper article can't be read:\n" + "\n".join(" - " + p for p in problems))
    # the masthead only heads the article, so a scenario that names one must have an article
    if d.MASTHEAD and not article:
        sys.exit(f"The matrix names a masthead ({d.MASTHEAD}) but no newspaper article was found in the packet "
                 f"({os.path.basename(packet)}). Check its '{ARTICLE_TAB}' tab, or remove the Masthead setting.")
    print(f"Newspaper article: {article['headline']}" if article else "No newspaper article in the packet")

    out = os.path.join(DOCS, sid)
    shutil.rmtree(out, ignore_errors=True)  # a clean build: nothing left over from an earlier one
    assets = os.path.join(folder, "assets")
    for w in d.WITNESSES:
        copy(find_file(assets, w["portrait"]), os.path.join(out, "assets", "portraits", w["portrait"]))
    copy(find_file(assets, d.MAP), os.path.join(out, "assets", d.MAP))
    for fn in (d.LOGO, d.MASTHEAD):
        if fn:
            own = os.path.join(assets, fn)
            copy(own if os.path.exists(own) else find_file(SHARED_ASSETS, fn), os.path.join(out, "assets", fn))
    for fn in ICONS:
        copy(os.path.join(HERE, fn), os.path.join(out, "assets", fn))

    data = json.dumps(scenario_json(d, article), ensure_ascii=False).replace("</", "<\\/")
    header = (
        f":: StoryTitle\n{d.TITLE}\n\n\n"
        ":: StoryData\n" + json.dumps({"ifid": d.IFID, "format": "SugarCube", "format-version": "2.37.3",
                                       "start": "Start"}, indent=2) + "\n\n\n"
    )
    text = json.dumps(g.text, ensure_ascii=False).replace("</", "<\\/")
    scen = (f":: ScenarioData [script]\n/* Generated by build.py from {os.path.basename(matrix)} and "
            f"{os.path.basename(GAME_TEXT)}. Do not edit by hand. */\nsetup.scenario = {data};\nsetup.gameText = {text};\n")
    with open(ENGINE, encoding="utf-8") as f:
        engine = f.read()
    # scenario data must load before the engine script, so it comes first
    twee = header + scen + "\n\n" + engine
    twee_path = os.path.join(out, "source", "story.twee")
    os.makedirs(os.path.dirname(twee_path), exist_ok=True)
    with open(twee_path, "w", encoding="utf-8") as f:
        f.write(twee)

    tweego = os.environ.get("TWEEGO") or shutil.which("tweego")
    if not tweego:
        sys.exit(f"Wrote {twee_path}, but Tweego was not found. Install it or set TWEEGO=/path/to/tweego, then rerun.")
    r = subprocess.run([tweego, "-f", "sugarcube-2", "--head", HEAD, "-o", os.path.join(out, "index.html"), twee_path],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit("Tweego failed:\n" + r.stderr)
    if r.stderr.strip():
        print(r.stderr.strip())
    print(f"Built {os.path.relpath(os.path.join(out, 'index.html'), ROOT)}")
    return {"id": sid, "title": d.TITLE, "pitch": d.PITCH or d.PUZZLE["mystery"], "ifid": d.IFID}


def build_landing(built):
    """docs/index.html: landing.html with one entry per scenario that has a built game in docs/."""
    import html
    entries = []
    for sid in all_scenarios():
        if not os.path.exists(os.path.join(DOCS, sid, "index.html")):
            print(f"Note: scenarios/{sid}/ has no built game yet, so the landing page leaves it out.")
            continue
        s = built.get(sid) or summary(sid)
        entries.append(f'<li class="scenario"><h3><a href="{html.escape(sid)}/index.html">{html.escape(s["title"])}</a></h3>'
                       f'<p>{html.escape(s["pitch"])}</p><p><a class="play" href="{html.escape(sid)}/index.html">Play</a></p></li>')
    with open(LANDING, encoding="utf-8") as f:
        page = f.read()
    if "{{SCENARIOS}}" not in page:
        sys.exit(f"{os.path.basename(LANDING)} has no {{{{SCENARIOS}}}} placeholder for the list of scenarios")
    with open(os.path.join(DOCS, "index.html"), "w", encoding="utf-8") as f:
        f.write(page.replace("{{SCENARIOS}}", "\n".join(entries)))
    for fn in ICONS:
        copy(os.path.join(HERE, fn), os.path.join(DOCS, "assets", fn))
    open(os.path.join(DOCS, ".nojekyll"), "w").close()  # serve the files as they are, without Jekyll
    print(f"Built docs/index.html ({len(entries)} scenario{'s' if len(entries) != 1 else ''})")


def summary(sid):
    """Title, summary and IFID of a scenario that is not being built, read without the full matrix checks."""
    import openpyxl
    matrix, _ = scenario_files(sid)
    wb = openpyxl.load_workbook(matrix, read_only=True, data_only=True)
    settings = {str(r[0]).strip(): str(r[1] or "").strip() for r in wb["Settings"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    pitch = ""
    for r in wb["Passages"].iter_rows(min_row=2, values_only=True):
        if r and str(r[0] or "").strip() == "Start" and str(r[1] or "").strip() == "Summary":
            pitch = str(r[2] or "").strip()
    return {"id": sid, "title": settings.get("Title", sid), "pitch": pitch, "ifid": settings.get("IFID", "")}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("scenarios", nargs="*", metavar="SCENARIO_ID", help="the scenarios to build (default: all of them)")
    ap.add_argument("--no-packet-check", action="store_true", help="build without checking the packets against the matrices")
    ap.add_argument("--framework", help="an export of the framework Google Doc (.docx, .txt or .pdf)")
    ap.add_argument("--no-framework-check", action="store_true", help="build without checking the game text against the framework")
    ap.add_argument("--downloads", metavar="FOLDER", help="a folder of downloads of the Google Sheets (.xlsx) and Docs "
                                                          "(.docx); each replaces its copy when the content differs")
    a = ap.parse_args()

    ids = a.scenarios or all_scenarios()
    if not ids:
        sys.exit("No scenarios in scenarios/")
    for sid in ids:
        scenario_files(sid)  # fail before any work if a name is wrong
    if a.downloads and not os.path.isdir(a.downloads):
        sys.exit(f"--downloads: no folder {a.downloads}")
    update_copies(a.downloads, ids)
    g = load_game_text(a.framework, a.no_framework_check, a.downloads)
    check_engine_keys(g)
    built = {sid: build_scenario(sid, g, a.no_packet_check) for sid in ids}

    # every scenario needs its own IFID, or players' saves would collide
    ifids = {}
    for sid in all_scenarios():
        ifid = (built.get(sid) or summary(sid))["ifid"].upper()
        if ifid in ifids:
            sys.exit(f"scenarios/{sid}/ and scenarios/{ifids[ifid]}/ have the same IFID; give each scenario its own.")
        ifids[ifid] = sid
    print()
    build_landing(built)


if __name__ == "__main__":
    main()
